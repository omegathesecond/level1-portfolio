#!/usr/bin/env python3
"""
Scraper pour récupérer la liste de tous les CFA de France
depuis lapprenti.com
"""

import requests
from bs4 import BeautifulSoup
import re
import time
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_URL = "https://www.lapprenti.com/html/regions/services_region.asp"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# Mapping région -> départements
REGIONS = {
    1: {"name": "Grand-Est", "deps": ["08", "10", "51", "52", "54", "55", "57", "67", "68", "88"]},
    2: {"name": "Nouvelle-Aquitaine", "deps": ["16", "17", "19", "23", "24", "33", "40", "47", "64", "79", "86", "87"]},
    3: {"name": "Occitanie", "deps": ["09", "11", "12", "30", "31", "32", "34", "46", "48", "65", "66", "81", "82"]},
    4: {"name": "Bourgogne-Franche-Comté", "deps": ["21", "25", "39", "58", "70", "71", "89", "90"]},
    5: {"name": "Normandie", "deps": ["14", "27", "50", "61", "76"]},
    6: {"name": "Auvergne-Rhône-Alpes", "deps": ["01", "03", "07", "15", "26", "38", "42", "43", "63", "69", "73", "74"]},
    7: {"name": "Bretagne", "deps": ["22", "29", "35", "56"]},
    8: {"name": "Ile-de-France", "deps": ["75", "77", "78", "91", "92", "93", "94", "95"]},
    9: {"name": "Hauts-de-France", "deps": ["02", "59", "60", "62", "80"]},
    10: {"name": "Centre-Val de Loire", "deps": ["18", "28", "36", "37", "41", "45"]},
    11: {"name": "Pays de la Loire", "deps": ["44", "49", "53", "72", "85"]},
    12: {"name": "Provence-Alpes-Côte d'Azur", "deps": ["04", "05", "06", "13", "83", "84"]},
    13: {"name": "La Réunion", "deps": ["974"]},
    14: {"name": "Guyane", "deps": ["973"]},
    15: {"name": "Martinique", "deps": ["972"]},
    16: {"name": "Guadeloupe", "deps": ["971"]},
    17: {"name": "Corse", "deps": ["2A", "2B"]},
}

# Noms des départements
DEPT_NAMES = {
    "01": "Ain", "02": "Aisne", "03": "Allier", "04": "Alpes-de-Haute-Provence",
    "05": "Hautes-Alpes", "06": "Alpes-Maritimes", "07": "Ardèche", "08": "Ardennes",
    "09": "Ariège", "10": "Aube", "11": "Aude", "12": "Aveyron",
    "13": "Bouches-du-Rhône", "14": "Calvados", "15": "Cantal", "16": "Charente",
    "17": "Charente-Maritime", "18": "Cher", "19": "Corrèze", "21": "Côte-d'Or",
    "22": "Côtes-d'Armor", "23": "Creuse", "24": "Dordogne", "25": "Doubs",
    "26": "Drôme", "27": "Eure", "28": "Eure-et-Loir", "29": "Finistère",
    "2A": "Corse-du-Sud", "2B": "Haute-Corse",
    "30": "Gard", "31": "Haute-Garonne", "32": "Gers", "33": "Gironde",
    "34": "Hérault", "35": "Ille-et-Vilaine", "36": "Indre", "37": "Indre-et-Loire",
    "38": "Isère", "39": "Jura", "40": "Landes", "41": "Loir-et-Cher",
    "42": "Loire", "43": "Haute-Loire", "44": "Loire-Atlantique", "45": "Loiret",
    "46": "Lot", "47": "Lot-et-Garonne", "48": "Lozère", "49": "Maine-et-Loire",
    "50": "Manche", "51": "Marne", "52": "Haute-Marne", "53": "Mayenne",
    "54": "Meurthe-et-Moselle", "55": "Meuse", "56": "Morbihan", "57": "Moselle",
    "58": "Nièvre", "59": "Nord", "60": "Oise", "61": "Orne",
    "62": "Pas-de-Calais", "63": "Puy-de-Dôme", "64": "Pyrénées-Atlantiques",
    "65": "Hautes-Pyrénées", "66": "Pyrénées-Orientales", "67": "Bas-Rhin",
    "68": "Haut-Rhin", "69": "Rhône", "70": "Haute-Saône", "71": "Saône-et-Loire",
    "72": "Sarthe", "73": "Savoie", "74": "Haute-Savoie", "75": "Paris",
    "76": "Seine-Maritime", "77": "Seine-et-Marne", "78": "Yvelines",
    "79": "Deux-Sèvres", "80": "Somme", "81": "Tarn", "82": "Tarn-et-Garonne",
    "83": "Var", "84": "Vaucluse", "85": "Vendée", "86": "Vienne",
    "87": "Haute-Vienne", "88": "Vosges", "89": "Yonne", "90": "Territoire de Belfort",
    "91": "Essonne", "92": "Hauts-de-Seine", "93": "Seine-Saint-Denis",
    "94": "Val-de-Marne", "95": "Val-d'Oise",
    "971": "Guadeloupe", "972": "Martinique", "973": "Guyane", "974": "La Réunion",
}


def fetch_page(url, retries=3):
    """Fetch a page with retry logic."""
    for attempt in range(retries):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=30)
            resp.encoding = resp.apparent_encoding
            return resp.text
        except Exception as e:
            if attempt < retries - 1:
                time.sleep(2 ** attempt)
            else:
                print(f"  ERREUR: impossible de charger {url}: {e}")
                return None


def parse_cfa_service(text):
    """Parse phone and address from cfa_service text."""
    phone = ""
    address = ""
    postal_code = ""
    city = ""

    if not text:
        return phone, address, postal_code, city

    # Format: "phone - address, postal_code city"
    # Phone patterns: 0X XX XX XX XX or 0X.XX.XX.XX.XX or 0XXXXXXXXX
    phone_match = re.match(
        r'(0[1-9][\s./-]?\d{2}[\s./-]?\d{2}[\s./-]?\d{2}[\s./-]?\d{2})\s*[-–]\s*(.*)',
        text
    )

    if phone_match:
        phone = phone_match.group(1).strip()
        addr_part = phone_match.group(2).strip()
    else:
        # Try alternate: phone might be after a dash or no phone
        addr_part = text.strip()

    # Extract postal code and city from the end
    pc_match = re.search(r',?\s*(\d{5})\s+(.+?)$', addr_part)
    if pc_match:
        postal_code = pc_match.group(1)
        city = pc_match.group(2).strip()
        address = addr_part[:pc_match.start()].strip().rstrip(',').strip()
    else:
        address = addr_part

    return phone, address, postal_code, city


def scrape_department(rg, dep_code):
    """Scrape all CFAs for a given department."""
    url = f"{BASE_URL}?action=2&rg={rg}&idDep={dep_code}"
    html = fetch_page(url)
    if not html:
        return []

    soup = BeautifulSoup(html, 'html.parser')
    cfas = []
    seen_names = set()

    # 1. Featured CFAs (cfaref-label sections)
    for section in soup.find_all('section', class_='cfaref-label'):
        link = section.find('a', class_='linkcfa')
        name = link.get_text(strip=True) if link else ""
        if not name:
            h2 = section.find('h2')
            name = h2.get_text(strip=True) if h2 else ""

        service = section.find('p', class_='cfa_service')
        service_text = service.get_text(strip=True) if service else ""
        phone, address, postal_code, city = parse_cfa_service(service_text)

        if name and name not in seen_names:
            seen_names.add(name)
            cfas.append({
                'name': name,
                'phone': phone,
                'address': address,
                'postal_code': postal_code,
                'city': city,
                'featured': True,
            })

    # 2. Regular CFAs (cfa_caption sections)
    for section in soup.find_all('section', class_='cfa_caption'):
        h2 = section.find('h2')
        if not h2:
            continue
        # Remove tooltip icon text
        for i_tag in h2.find_all('i'):
            i_tag.decompose()
        name = h2.get_text(strip=True)

        service = section.find('p', class_='cfa_service')
        service_text = service.get_text(strip=True) if service else ""
        phone, address, postal_code, city = parse_cfa_service(service_text)

        if name and name not in seen_names:
            seen_names.add(name)
            cfas.append({
                'name': name,
                'phone': phone,
                'address': address,
                'postal_code': postal_code,
                'city': city,
                'featured': False,
            })

    return cfas


def create_excel(all_data, output_path):
    """Create an Excel file with all CFA data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Liste des CFA de France"

    # Styles
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Headers
    headers = [
        "N°", "Nom du CFA", "Téléphone", "Adresse",
        "Code Postal", "Ville", "N° Département", "Département",
        "Région", "Référencé"
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data
    row = 2
    for entry in all_data:
        ws.cell(row=row, column=1, value=row - 1).border = thin_border
        ws.cell(row=row, column=2, value=entry['name']).border = thin_border
        ws.cell(row=row, column=3, value=entry['phone']).border = thin_border
        ws.cell(row=row, column=4, value=entry['address']).border = thin_border
        ws.cell(row=row, column=5, value=entry['postal_code']).border = thin_border
        ws.cell(row=row, column=6, value=entry['city']).border = thin_border
        ws.cell(row=row, column=7, value=entry['dep_code']).border = thin_border
        ws.cell(row=row, column=8, value=entry['dep_name']).border = thin_border
        ws.cell(row=row, column=9, value=entry['region']).border = thin_border
        ws.cell(row=row, column=10, value="Oui" if entry['featured'] else "Non").border = thin_border
        row += 1

    # Adjust column widths
    col_widths = [6, 55, 18, 45, 12, 25, 15, 25, 30, 12]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = width

    # Freeze first row
    ws.freeze_panes = 'A2'

    # Auto-filter
    ws.auto_filter.ref = f"A1:J{row - 1}"

    # Add summary sheet
    ws2 = wb.create_sheet("Résumé par Région")
    ws2.cell(row=1, column=1, value="Région").font = header_font
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=2, value="Nombre de CFA").font = header_font
    ws2.cell(row=1, column=2).fill = header_fill
    ws2.column_dimensions['A'].width = 35
    ws2.column_dimensions['B'].width = 18

    # Count by region
    region_counts = {}
    for entry in all_data:
        region_counts[entry['region']] = region_counts.get(entry['region'], 0) + 1

    r = 2
    for region, count in sorted(region_counts.items()):
        ws2.cell(row=r, column=1, value=region)
        ws2.cell(row=r, column=2, value=count)
        r += 1
    ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    ws2.cell(row=r, column=2, value=len(all_data)).font = Font(bold=True)

    wb.save(output_path)
    print(f"\nFichier Excel créé: {output_path}")
    print(f"Total CFA: {len(all_data)}")


def main():
    all_cfas = []
    total_regions = len(REGIONS)

    for rg_idx, (rg, info) in enumerate(REGIONS.items(), 1):
        region_name = info['name']
        deps = info['deps']
        print(f"\n[{rg_idx}/{total_regions}] Région: {region_name} ({len(deps)} départements)")

        for dep_code in deps:
            dep_name = DEPT_NAMES.get(dep_code, dep_code)
            print(f"  Département {dep_code} - {dep_name}...", end=" ", flush=True)

            cfas = scrape_department(rg, dep_code)
            for cfa in cfas:
                cfa['dep_code'] = dep_code
                cfa['dep_name'] = dep_name
                cfa['region'] = region_name

            all_cfas.extend(cfas)
            print(f"{len(cfas)} CFA trouvés")

            # Be polite - small delay between requests
            time.sleep(0.5)

    print(f"\n{'='*60}")
    print(f"TOTAL: {len(all_cfas)} CFA récupérés sur toute la France")
    print(f"{'='*60}")

    output_path = "/home/user/level1-portfolio/Liste_CFA_France.xlsx"
    create_excel(all_cfas, output_path)

    # Also save as JSON for reference
    json_path = "/home/user/level1-portfolio/liste_cfa_france.json"
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(all_cfas, f, ensure_ascii=False, indent=2)
    print(f"Fichier JSON créé: {json_path}")


if __name__ == "__main__":
    main()
