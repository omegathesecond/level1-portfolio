#!/usr/bin/env python3
"""
Crée un fichier Excel combinant les CFAs de 3 sources :
  - data.gouv.fr IDF  (cfa_idf_datagouv.csv)
  - lapprenti.com     (liste_cfa_france.json)
  - letudiant.fr      (liste_cfa_letudiant.json)

Enrichit chaque CFA avec son site web et email depuis :
  - annuaire_education_apprentissage.json
"""

import json
import unicodedata
import re
import csv
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def normalize(text: str) -> str:
    """Normalise un texte pour la comparaison : minuscules, sans accents, sans ponctuation."""
    if not text:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = text.lower()
    text = re.sub(r"[^a-z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def dep_from_cp(code_postal: str) -> tuple:
    """Retourne (dep_code, dep_name) depuis un code postal."""
    DEP_NAMES = {
        "01": "Ain", "02": "Aisne", "03": "Allier", "04": "Alpes-de-Haute-Provence",
        "05": "Hautes-Alpes", "06": "Alpes-Maritimes", "07": "Ardèche", "08": "Ardennes",
        "09": "Ariège", "10": "Aube", "11": "Aude", "12": "Aveyron",
        "13": "Bouches-du-Rhône", "14": "Calvados", "15": "Cantal", "16": "Charente",
        "17": "Charente-Maritime", "18": "Cher", "19": "Corrèze", "2A": "Corse-du-Sud",
        "2B": "Haute-Corse", "21": "Côte-d'Or", "22": "Côtes-d'Armor", "23": "Creuse",
        "24": "Dordogne", "25": "Doubs", "26": "Drôme", "27": "Eure",
        "28": "Eure-et-Loir", "29": "Finistère", "30": "Gard", "31": "Haute-Garonne",
        "32": "Gers", "33": "Gironde", "34": "Hérault", "35": "Ille-et-Vilaine",
        "36": "Indre", "37": "Indre-et-Loire", "38": "Isère", "39": "Jura",
        "40": "Landes", "41": "Loir-et-Cher", "42": "Loire", "43": "Haute-Loire",
        "44": "Loire-Atlantique", "45": "Loiret", "46": "Lot", "47": "Lot-et-Garonne",
        "48": "Lozère", "49": "Maine-et-Loire", "50": "Manche", "51": "Marne",
        "52": "Haute-Marne", "53": "Mayenne", "54": "Meurthe-et-Moselle", "55": "Meuse",
        "56": "Morbihan", "57": "Moselle", "58": "Nièvre", "59": "Nord",
        "60": "Oise", "61": "Orne", "62": "Pas-de-Calais", "63": "Puy-de-Dôme",
        "64": "Pyrénées-Atlantiques", "65": "Hautes-Pyrénées", "66": "Pyrénées-Orientales",
        "67": "Bas-Rhin", "68": "Haut-Rhin", "69": "Rhône", "70": "Haute-Saône",
        "71": "Saône-et-Loire", "72": "Sarthe", "73": "Savoie", "74": "Haute-Savoie",
        "75": "Paris", "76": "Seine-Maritime", "77": "Seine-et-Marne",
        "78": "Yvelines", "79": "Deux-Sèvres", "80": "Somme", "81": "Tarn",
        "82": "Tarn-et-Garonne", "83": "Var", "84": "Vaucluse", "85": "Vendée",
        "86": "Vienne", "87": "Haute-Vienne", "88": "Vosges", "89": "Yonne",
        "90": "Territoire de Belfort", "91": "Essonne", "92": "Hauts-de-Seine",
        "93": "Seine-Saint-Denis", "94": "Val-de-Marne", "95": "Val-d'Oise",
        "971": "Guadeloupe", "972": "Martinique", "973": "Guyane", "974": "La Réunion",
        "976": "Mayotte",
    }
    if not code_postal:
        return ("", "")
    cp = code_postal.strip()
    if cp.startswith("97"):
        code = cp[:3]
    elif cp.startswith("20"):
        code = "2A"  # approximation
    else:
        code = cp[:2]
    return (code, DEP_NAMES.get(code, ""))


# ---------------------------------------------------------------------------
# Chargement de l'annuaire (référentiel web + email)
# ---------------------------------------------------------------------------

def load_annuaire() -> tuple:
    """
    Retourne deux dicts :
      rne_lookup  : { rne_code → {"web": ..., "mail": ..., "telephone": ...} }
      name_lookup : { (nom_normalisé, code_postal) → {"web": ..., "mail": ...} }
    """
    path = os.path.join(BASE_DIR, "annuaire_education_apprentissage.json")
    print("Chargement annuaire éducation…")
    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    rne_lookup = {}
    name_lookup = {}

    for item in data:
        web = item.get("web", "") or ""
        mail = item.get("mail", "") or ""
        tel = item.get("telephone", "") or ""
        info = {"web": web.strip(), "mail": mail.strip(), "telephone": tel.strip()}

        rne = item.get("identifiant_de_l_etablissement", "")
        if rne:
            rne_lookup[rne.strip().upper()] = info

        nom = normalize(item.get("nom_etablissement", ""))
        cp = (item.get("code_postal", "") or "").strip()
        if nom and cp:
            key = (nom, cp)
            if key not in name_lookup:
                name_lookup[key] = info

    print(f"  → {len(rne_lookup)} entrées par RNE, {len(name_lookup)} entrées par nom+CP")
    return rne_lookup, name_lookup


def lookup_web(name: str, cp: str, rne: str,
               rne_lookup: dict, name_lookup: dict) -> dict:
    """Cherche les infos web/mail pour un CFA."""
    # 1) Par RNE
    if rne:
        info = rne_lookup.get(rne.strip().upper())
        if info:
            return info
    # 2) Par nom normalisé + code postal
    key = (normalize(name), cp.strip())
    info = name_lookup.get(key)
    if info:
        return info
    return {"web": "", "mail": "", "telephone": ""}


# ---------------------------------------------------------------------------
# Chargement des 3 sources
# ---------------------------------------------------------------------------

def load_datagouv_idf(rne_lookup, name_lookup) -> list:
    path = os.path.join(BASE_DIR, "cfa_idf_datagouv.csv")
    print("Chargement data.gouv.fr IDF…")
    rows = []
    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for r in reader:
            name = (r.get("nom_de_l_organisme") or r.get("nom_cfa_conventionnel") or "").strip()
            adresse = (r.get("adresse_1") or r.get("adresse_cfa_conventionnel") or "").strip()
            cp = (r.get("code_postal") or r.get("code_postal_cfa_conventionnel") or "").strip()
            ville = (r.get("commune") or r.get("ville_cfa_conventionnel") or "").strip()
            dep = (r.get("departement") or "").strip()
            rne = (r.get("code_rne") or "").strip()

            dep_code, dep_name = dep_from_cp(cp)
            if not dep:
                dep = dep_name

            info = lookup_web(name, cp, rne, rne_lookup, name_lookup)
            rows.append({
                "Source": "data.gouv.fr IDF",
                "Nom du CFA": name,
                "Adresse": adresse,
                "Code Postal": cp,
                "Ville": ville,
                "Département": dep,
                "Région": "Île-de-France",
                "Téléphone": info.get("telephone", ""),
                "Site Web": info.get("web", ""),
                "Email": info.get("mail", ""),
            })
    print(f"  → {len(rows)} CFAs chargés")
    return rows


def load_lapprenti(rne_lookup, name_lookup) -> list:
    path = os.path.join(BASE_DIR, "liste_cfa_france.json")
    print("Chargement lapprenti.com…")
    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    rows = []
    for item in data:
        name = (item.get("name") or "").strip()
        cp = (item.get("postal_code") or "").strip()
        info = lookup_web(name, cp, "", rne_lookup, name_lookup)
        rows.append({
            "Source": "lapprenti.com",
            "Nom du CFA": name,
            "Adresse": (item.get("address") or "").strip(),
            "Code Postal": cp,
            "Ville": (item.get("city") or "").strip(),
            "Département": (item.get("dep_name") or "").strip(),
            "Région": (item.get("region") or "").strip(),
            "Téléphone": (item.get("phone") or "").strip(),
            "Site Web": info.get("web", ""),
            "Email": info.get("mail", ""),
        })
    print(f"  → {len(rows)} CFAs chargés")
    return rows


def load_letudiant(rne_lookup, name_lookup) -> list:
    path = os.path.join(BASE_DIR, "liste_cfa_letudiant.json")
    print("Chargement letudiant.fr…")
    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    rows = []
    for item in data:
        name = (item.get("name") or "").strip()
        cp = (item.get("postal_code") or "").strip()
        city = (item.get("city_full") or item.get("city") or "").strip()
        dep_code = (item.get("dep_code") or "").strip()
        dep_name = (item.get("dep_name") or "").strip()
        region = (item.get("region") or "").strip()

        # Si pas de département, dériver depuis le CP
        if not dep_code and cp:
            dep_code, dep_name = dep_from_cp(cp)

        info = lookup_web(name, cp, "", rne_lookup, name_lookup)
        rows.append({
            "Source": "letudiant.fr",
            "Nom du CFA": name,
            "Adresse": (item.get("address") or "").strip(),
            "Code Postal": cp,
            "Ville": city,
            "Département": dep_name,
            "Région": region,
            "Téléphone": (item.get("phone") or "").strip(),
            "Site Web": info.get("web", ""),
            "Email": info.get("mail", ""),
        })
    print(f"  → {len(rows)} CFAs chargés")
    return rows


# ---------------------------------------------------------------------------
# Génération Excel
# ---------------------------------------------------------------------------

COLUMNS = [
    "Source", "Nom du CFA", "Adresse", "Code Postal",
    "Ville", "Département", "Région", "Téléphone", "Site Web", "Email",
]

COL_WIDTHS = {
    "Source": 20,
    "Nom du CFA": 45,
    "Adresse": 40,
    "Code Postal": 12,
    "Ville": 25,
    "Département": 25,
    "Région": 22,
    "Téléphone": 16,
    "Site Web": 40,
    "Email": 35,
}

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

SOURCE_COLORS = {
    "data.gouv.fr IDF": "D6E4F0",
    "lapprenti.com":    "D5F5E3",
    "letudiant.fr":     "FEF9E7",
}


def create_excel(all_rows: list, out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tous les CFA"

    # En-têtes
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # Données
    for row_idx, row in enumerate(all_rows, start=2):
        source = row.get("Source", "")
        fill_color = SOURCE_COLORS.get(source, "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center")

    # Largeurs des colonnes
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 20)

    # Filtre automatique
    ws.auto_filter.ref = ws.dimensions

    wb.save(out_path)
    print(f"\nFichier Excel créé : {out_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    rne_lookup, name_lookup = load_annuaire()

    rows_datagouv = load_datagouv_idf(rne_lookup, name_lookup)
    rows_lapprenti = load_lapprenti(rne_lookup, name_lookup)
    rows_letudiant = load_letudiant(rne_lookup, name_lookup)

    all_rows = rows_datagouv + rows_lapprenti + rows_letudiant
    total = len(all_rows)

    with_web = sum(1 for r in all_rows if r.get("Site Web"))
    with_email = sum(1 for r in all_rows if r.get("Email"))

    print(f"\n{'='*55}")
    print(f"  TOTAL CFAs combinés : {total}")
    print(f"  - data.gouv.fr IDF  : {len(rows_datagouv)}")
    print(f"  - lapprenti.com     : {len(rows_lapprenti)}")
    print(f"  - letudiant.fr      : {len(rows_letudiant)}")
    print(f"  CFAs avec site web  : {with_web} ({with_web*100//total}%)")
    print(f"  CFAs avec email     : {with_email} ({with_email*100//total}%)")
    print(f"{'='*55}")

    out_path = os.path.join(BASE_DIR, "Liste_CFA_Combinée.xlsx")
    create_excel(all_rows, out_path)


if __name__ == "__main__":
    main()
